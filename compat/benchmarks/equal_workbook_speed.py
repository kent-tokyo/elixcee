"""Paired durable XLSX benchmark. Generated outputs are verified, never published."""
import argparse
import hashlib
import json
import math
import os
import platform
import statistics
import subprocess
import sys
import tempfile
import time
from pathlib import Path

import openpyxl


def durable_sync(file):
    # Match this Rust toolchain's std::fs::File::sync_all implementation.
    # On macOS os.fsync alone does NOT match F_FULLFSYNC's durability barrier.
    if sys.platform == 'darwin':
        import fcntl
        fcntl.fcntl(file.fileno(), fcntl.F_FULLFSYNC)
    else:
        os.fsync(file.fileno())


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def logical(path):
    wb = openpyxl.load_workbook(path, data_only=False)
    result = {ws.title: {cell.coordinate: cell.value for row in ws for cell in row
                         if cell.value is not None} for ws in wb}
    wb.close()
    return result


def check(path, expected):
    assert logical(path) == expected, f"value/formula mismatch: {path}"


def py_batch(fixture, output, iterations, expected):
    samples = []
    for _ in range(iterations):
        start = time.perf_counter()
        wb = openpyxl.load_workbook(fixture, data_only=False)
        loaded = (time.perf_counter() - start) * 1000
        ws = wb.worksheets[0]
        ws['A1'] = 123
        ws['B1'] = '=1+2'
        mutated = (time.perf_counter() - start) * 1000
        # Same filesystem and publication order as elixcee: create beside output,
        # finish ZIP, sync_all-equivalent barrier, close, then rename.
        # Neither syncs the directory.
        with tempfile.NamedTemporaryFile(dir=output.parent, suffix='.xlsx', delete=False) as f:
            temporary = Path(f.name)
            wb.save(f)
            f.flush()
            durable_sync(f)
        os.replace(temporary, output)
        saved = (time.perf_counter() - start) * 1000
        reread = openpyxl.load_workbook(output, data_only=False)
        elapsed = (time.perf_counter() - start) * 1000
        values = {ws.title: {c.coordinate: c.value for row in ws for c in row
                            if c.value is not None} for ws in reread}
        assert values == expected
        reread.close()
        wb.close()
        samples.append(dict(load_ms=loaded, mutate_ms=mutated-loaded,
                            save_ms=saved-mutated, reload_ms=elapsed-saved, total_ms=elapsed))
    return samples


def main():
    p = argparse.ArgumentParser()
    p.add_argument('--before', type=Path, help='optional pre-change binary for causal comparison')
    p.add_argument('--after', type=Path, required=True)
    p.add_argument('--output', type=Path, required=True)
    p.add_argument('--rounds', type=int, default=3)
    p.add_argument('--iterations', type=int, default=10)
    args = p.parse_args()
    assert args.rounds >= 3 and args.iterations >= 5
    repo = Path(__file__).resolve().parents[2]
    with tempfile.TemporaryDirectory(prefix='elixcee-equal-bench-') as directory:
        directory = Path(directory)
        fixtures = [('small', repo/'tests/fixtures/e2e/source.xlsx')]
        for rows in (1000, 10000):
            path = directory/f'numeric-{rows}x10.xlsx'
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = 'Sheet1'
            for row in range(1, rows+1):
                ws.append([row * 10 + col for col in range(1, 11)])
            wb.save(path)
            wb.close()
            fixtures.append((f'numeric-{rows}x10', path))
        results = []
        binaries = {'after': args.after.resolve()}
        if args.before:
            binaries = {'before': args.before.resolve(), **binaries}
        for name, fixture in fixtures:
            expected = logical(fixture)
            assert len(expected) == 1
            sheet = next(iter(expected.values()))
            sheet.update(A1=123, B1='=1+2')
            output = directory/'result.xlsx'
            observations = {key: [] for key in (*binaries, 'openpyxl')}
            sizes = {}
            def run(key, count):
                if key == 'openpyxl':
                    samples = py_batch(fixture, output, count, expected)
                else:
                    r = subprocess.run([str(binaries[key]), str(fixture), str(output), str(count)],
                                       capture_output=True, text=True, check=True, timeout=180)
                    payload = json.loads(r.stdout)
                    samples = payload['samples']
                check(output, expected)
                sizes[key] = output.stat().st_size
                return samples
            for key in observations:
                run(key, 2)  # Explicitly excluded warmup; OS cache is warm for all.
            orders = []
            for round_no in range(args.rounds):
                keys = list(observations)
                offset = round_no % len(keys)
                keys = keys[offset:] + keys[:offset]
                orders.append(keys)
                for key in keys:
                    samples = run(key, args.iterations)
                    observations[key].extend(dict(round=round_no+1, **s) for s in samples)
            summaries = {}
            for key, samples in observations.items():
                summaries[key] = {}
                for metric in ('load_ms','mutate_ms','save_ms','reload_ms','total_ms'):
                    values = sorted(s[metric] for s in samples)
                    summaries[key][metric] = dict(p50=statistics.median(values),
                        p95=values[math.ceil(.95*len(values))-1])
            result = dict(fixture=name, fixture_sha256=digest(fixture),
                          input_bytes=fixture.stat().st_size, output_bytes=sizes,
                          verified=True, order=orders, summaries=summaries, samples=observations)
            results.append(result)
            print(name, json.dumps(summaries), flush=True)
        payload = dict(schema_version=1, date=time.strftime('%Y-%m-%d'),
            platform=platform.platform(), machine=platform.machine(), python=platform.python_version(),
            openpyxl=openpyxl.__version__, elixcee='1.0.1-unreleased-working-tree',
            openpyxl_lxml=openpyxl.LXML,
            rustc=subprocess.check_output(['rustc','--version'],text=True).strip(),
            head=subprocess.check_output(['git','rev-parse','HEAD'],cwd=repo,text=True).strip(),
            binary_sha256={k:digest(v) for k,v in binaries.items()},
            source_sha256={p:digest(repo/p) for p in (
                'src/lib.rs', 'src/reader.rs', 'src/vm/mod.rs', 'Cargo.lock',
                'examples/bench_workbook.rs', 'compat/benchmarks/equal_workbook_speed.py')},
            rounds=args.rounds, iterations_per_round=args.iterations,
            timing='load + mutate A1/B1 + save + sync_all-equivalent file barrier + atomic rename + reload; startup and validation excluded',
            file_sync='F_FULLFSYNC' if sys.platform == 'darwin' else 'fsync',
            validation='all cell values and formula strings via openpyxl; no formula evaluation or full OOXML equivalence claim',
            fixtures=results)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(json.dumps(payload, indent=2)+'\n')


if __name__ == '__main__':
    main()
