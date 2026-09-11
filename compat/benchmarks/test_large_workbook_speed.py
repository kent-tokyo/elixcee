import tempfile
import unittest
from io import BytesIO
from pathlib import Path

import openpyxl
from large_workbook_speed import check_output, make_fixture


class LargeWorkbookValidationTests(unittest.TestCase):
    def test_all_sheets_are_checked_and_only_first_sheet_is_edited(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'fixture.xlsx'
            make_fixture(path, 3, 4)
            wb = openpyxl.load_workbook(BytesIO(path.read_bytes()))
            wb.worksheets[0]['A1'] = 123
            wb.worksheets[0]['B1'] = '=1+2'
            wb.save(path)
            check_output(path, 3, 4)
            wb.worksheets[3]['J3'] = -1
            wb.save(path)
            with self.assertRaises(AssertionError):
                check_output(path, 3, 4)
            wb.close()

    def test_missing_extra_cells_rows_sheets_and_formula_loss_are_rejected(self):
        changes = [lambda wb: setattr(wb.worksheets[0]['B1'], 'value', 3),
                   lambda wb: setattr(wb.worksheets[1]['J3'], 'value', None),
                   lambda wb: setattr(wb.worksheets[1]['K3'], 'value', 1),
                   lambda wb: setattr(wb.worksheets[1]['A4'], 'value', 1),
                   lambda wb: wb.remove(wb.worksheets[1]),
                   lambda wb: wb.create_sheet('extra')]
        for change in changes:
            with self.subTest(change=change), tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / 'fixture.xlsx'
                make_fixture(path, 3, 2)
                wb = openpyxl.load_workbook(BytesIO(path.read_bytes()))
                wb.worksheets[0]['A1'] = 123
                wb.worksheets[0]['B1'] = '=1+2'
                change(wb)
                wb.save(path)
                wb.close()
                with self.assertRaises(AssertionError):
                    check_output(path, 3, 2)


if __name__ == '__main__':
    unittest.main()
