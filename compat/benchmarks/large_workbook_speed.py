"""Large-workbook paired benchmark; never relaxes the reader's XML budgets."""
import argparse
import gc
import itertools
import json
import os
import platform
import statistics
import subprocess
import tempfile
import time
from pathlib import Path

import openpyxl
from equal_workbook_speed import digest
from compare_closedxml import summarize
from paired_speedup import compare_parts


CASES = {'100k': (10000, 1), '400k': (40000, 1), '1m': (25000, 4)}


def make_fixture(path, rows, sheets):
    wb = openpyxl.Workbook(write_only=True)
    for index in range(sheets):
        ws = wb.create_sheet(f'Sheet{index + 1}')
        for row in range(1, rows + 1):
            ws.append([index * rows * 10 + row * 10 + col for col in range(1, 11)])
    wb.save(path)
    wb.close()


def check_output(path, rows, sheets):
    # Stream independent validation so 1m-cell Python models do not remain in
    # memory while Rust is timed. Detect extra/missing rows, cells and sheets.
    source = path.open('rb')
    wb = None
    try:
        wb = openpyxl.load_workbook(source, read_only=True, data_only=False)
        assert wb.sheetnames == [f'Sheet{i + 1}' for i in range(sheets)]
        for index, ws in enumerate(wb):
            expected_rows = range(1, rows + 1)
            for row, values in itertools.zip_longest(expected_rows, ws.values):
                assert row is not None and values is not None, 'row count changed'
                expected = [index * rows * 10 + row * 10 + col for col in range(1, 11)]
                if index == 0 and row == 1:
                    expected[:2] = [123, '=1+2']
                assert tuple(expected) == values, f'cell/formula changed: {index}:{row}'
    finally:
        if wb is not None:
            wb.close()
        source.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--before', type=Path, required=True)
    parser.add_argument('--after', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--pairs', type=int, default=20)
    parser.add_argument('--cases', choices=CASES, nargs='+', default=list(CASES))
    args = parser.parse_args()
    if args.pairs < 4 or args.pairs % 2:
        parser.error('pairs must be even and at least 4 (use >=20 for confirmation)')
    repo = Path(__file__).resolve().parents[2]
    binaries = {key: path.resolve(strict=True) for key, path in
                [('before', args.before), ('after', args.after)]}
    source_paths = ['src/lib.rs', 'src/reader.rs', 'src/vm/mod.rs', 'Cargo.toml',
                    'Cargo.lock', 'examples/bench_workbook.rs',
                    'compat/benchmarks/large_workbook_speed.py',
                    'compat/benchmarks/paired_speedup.py',
                    'compat/benchmarks/equal_workbook_speed.py',
                    'compat/benchmarks/compare_closedxml.py']
    metadata = dict(date=time.strftime('%Y-%m-%d'), platform=platform.platform(),
        python=platform.python_version(), openpyxl=openpyxl.__version__,
        rustc=subprocess.check_output(['rustc', '--version'], text=True).strip(),
        binary_sha256={key: digest(path) for key, path in binaries.items()},
        after_source_sha256={p: digest(repo / p) for p in source_paths},
        pairs=args.pairs, warmups_per_engine_per_fixture=3, target_speedup=1.2,
        confirmation=args.pairs >= 20,
        timing='load/edit first sheet A1/B1/standard durable save/rename/reload; startup, preload, assertions and destruction excluded',
        validation='all decompressed ZIP members byte-identical each pair; independent streaming openpyxl all-cell/formula checks each ten pairs and final',
        safety='unchanged XML/ZIP budgets; 1m cells split across four 250k-cell sheets',
        host_load_start=list(os.getloadavg()), fixtures=[])
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='elixcee-large-pairs-') as directory:
        directory = Path(directory)
        for case in args.cases:
            rows, sheets = CASES[case]
            fixture = directory / f'{case}.xlsx'
            make_fixture(fixture, rows, sheets)
            gc.collect()
            outputs = {key: directory / f'{key}.xlsx' for key in binaries}
            samples = {key: [] for key in binaries}
            versions = {}

            def run(key):
                result = subprocess.run([str(binaries[key]), str(fixture), str(outputs[key]), '1'],
                    capture_output=True, text=True, check=True, timeout=180)
                payload = json.loads(result.stdout)
                assert len(payload['samples']) == 1
                versions[key] = payload['version']
                return payload['samples'][0]

            for _ in range(3):
                for key in binaries:
                    run(key)
            orders = []
            for pair in range(args.pairs):
                order = ['before', 'after'] if pair % 2 == 0 else ['after', 'before']
                orders.append(order)
                for key in order:
                    samples[key].append(dict(pair=pair + 1, **run(key)))
                compare_parts(outputs['before'], outputs['after'])
                if (pair + 1) % 10 == 0 or pair + 1 == args.pairs:
                    for output in outputs.values():
                        check_output(output, rows, sheets)
                    gc.collect()
                print(f'{case}: pair {pair + 1}/{args.pairs}', flush=True)
            summary = {key: summarize(values) for key, values in samples.items()}
            ratio = summary['before']['total_ms']['p50'] / summary['after']['total_ms']['p50']
            result = dict(fixture=case, rows_per_sheet=rows, columns=10, sheets=sheets,
                cells=rows * 10 * sheets, fixture_sha256=digest(fixture), versions=versions,
                summaries=summary, p50_speedup=ratio,
                median_paired_speedup=statistics.median(
                    b['total_ms'] / a['total_ms'] for b, a in zip(samples['before'], samples['after'])),
                target_met=ratio >= 1.2, zip_parts_verified=True,
                output_bytes={key: path.stat().st_size for key, path in outputs.items()},
                orders=orders, samples=samples)
            metadata['fixtures'].append(result)
            metadata['host_load_end'] = list(os.getloadavg())
            args.output.write_text(json.dumps(metadata, indent=2) + '\n')
            print(case, json.dumps(summary), 'speedup', ratio, flush=True)


if __name__ == '__main__':
    main()
